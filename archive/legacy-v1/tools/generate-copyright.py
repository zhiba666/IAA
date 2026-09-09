"""Build copyright technical-material drafts using bundled Python dependencies.
Source export requires --frozen after the owner confirms source-tree readiness.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs/copyright'
OUT = ROOT / 'deliverables/copyright'
NAME, VERSION = '小小爆米花厂', '1.0.0'
STEM = f'{NAME}_{VERSION}'
# Runtime code, build/audio generation and the local-preview entry/server only.
SOURCE_FILES = ['src/main.js', 'src/core.js', 'src/platform.js', 'src/renderer.js',
                'src/audio.js', 'tools/build.mjs', 'tools/audio.mjs',
                'tools/serve.mjs', 'web/index.html']
SOURCES_URL = 'https://banshi.beijing.gov.cn/pubtask/task/1/110000000000/3e283672-76be-4c8c-98e8-0bebe9bd06bf_app.html'

def sha(data):
    return hashlib.sha256(data).hexdigest()

def set_font(style, size, bold=False, before=0, after=0, line=1.25):
    style.font.name='Calibri';style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False
    style.font.color.rgb=RGBColor.from_string('000000')
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    pf=style.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line

def field(p, code):
    item=OxmlElement('w:fldSimple');item.set(qn('w:instr'),code);p._p.append(item)

def make_manual(screenshots):
    doc=Document();props=doc.core_properties
    props.title=NAME+' '+VERSION+' 操作说明书';props.subject='内部草稿 / 待权利人确认'
    props.author='';props.last_modified_by='';props.keywords=''
    props.comments='Technical draft only; applicant identity and legal facts require confirmation.'
    sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=Mm(23);sec.left_margin=sec.right_margin=Mm(25)
    sec.header_distance=sec.footer_distance=Mm(12.5)
    # compact_reference_guide plus statutory-A4, Chinese-monochrome and
    # registration-density (body after=4pt) overrides; no font-size reduction.
    set_font(doc.styles['Normal'],11,after=4)
    set_font(doc.styles['Title'],25,True,after=10)
    set_font(doc.styles['Subtitle'],15,after=14)
    set_font(doc.styles['Heading 1'],16,True,before=18,after=10)
    set_font(doc.styles['Heading 2'],13,True,before=14,after=7)
    set_font(doc.styles['Heading 3'],12,True,before=10,after=5)
    set_font(doc.styles['Caption'],9,after=8,line=1.1)
    set_font(doc.styles['Header'],9,line=1);set_font(doc.styles['Footer'],9,line=1)
    # The bundled blank DOCX can carry inherited title rules and header tabs.
    # A formal monochrome application draft has no decorative rule or center tab.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for style_name in ['Header','Footer','Normal','Title','Subtitle']:
        doc.styles[style_name].paragraph_format.tab_stops.clear_all()
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].paragraph_format.keep_with_next=True
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        doc.styles[name].paragraph_format.widow_control=True
    h=sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.clear_all()
    h.paragraph_format.tab_stops.add_tab_stop(Mm(160),WD_TAB_ALIGNMENT.RIGHT)
    h.add_run(f'{NAME} {VERSION}  操作说明书\t');field(h,'PAGE')
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f.add_run('内部草稿 / 待权利人确认')
    content=(DOCS/'manual-draft.md').read_text(encoding='utf-8');on_cover=True
    for block in content.strip().split('\n\n'):
        block=block.strip()
        if block=='---':doc.add_page_break();on_cover=False;continue
        if block.startswith('# '):doc.add_paragraph(block[2:],'Title' if on_cover else 'Heading 1')
        elif block.startswith('## '):doc.add_paragraph(block[3:],'Subtitle' if on_cover else 'Heading 2')
        else:
            for line in block.splitlines():doc.add_paragraph(line)
    for index in range(0,len(screenshots),2):
        doc.add_page_break();doc.add_paragraph(f'附录：界面示例 {index+1}-{min(index+2,len(screenshots))}','Heading 1')
        if index==0:
            doc.add_paragraph('以下为当前源码渲染的操作界面示意，使用固定示例进度；图示不作为真机测试或广告验证证据。')
        pair=screenshots[index:index+2];images=doc.add_paragraph();images.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for j,item in enumerate(pair):
            if j:images.add_run('  ')
            shape=images.add_run().add_picture(str(Path(item['path']).resolve()),width=Mm(73))
            shape._inline.docPr.set('descr',item['caption'])
        for j,item in enumerate(pair):
            doc.add_paragraph(f'图{index+j+1}  {item["caption"]}','Caption')
            if item.get('description'):doc.add_paragraph(item['description'])
        doc.add_paragraph('图示为当前项目实际运行画面；数值和进度仅为操作示例。','Caption')
    target=OUT/(STEM+'_操作说明书_内部草稿.docx');doc.save(target);return target

def setup_pdf_fonts():
    pdfmetrics.registerFont(TTFont('CJK','C:/Windows/Fonts/simsun.ttc',subfontIndex=0))
    pdfmetrics.registerFont(TTFont('Mono','C:/Windows/Fonts/consola.ttf'))

def text_width(s,size):
    return sum(pdfmetrics.stringWidth(ch,'Mono' if ord(ch)<128 else 'CJK',size) for ch in s)

def draw_code(c,text,x,y,size):
    for chunk in re.findall(r'[\x00-\x7f]+|[^\x00-\x7f]+',text):
        font='Mono' if ord(chunk[0])<128 else 'CJK'
        c.setFont(font,size);c.drawString(x,y,chunk)
        x+=pdfmetrics.stringWidth(chunk,font,size)

def fold_line(s,width,size):
    # Concatenating all segments exactly recovers the tab-expanded source line.
    s=s.expandtabs(4)
    if not s:return ['']
    result=[]
    while s:
        used=0;end=0
        for ch in s:
            w=text_width(ch,size)
            if used+w>width:break
            used+=w;end+=1
        end=max(1,end)
        if end<len(s):
            breaks=[m.end() for m in re.finditer(r'[ ;,{}]',s[:end])]
            if breaks and breaks[-1]>end*.65:end=breaks[-1]
        result.append(s[:end]);s=s[end:]
    return result

def make_sources():
    setup_pdf_fonts();export=OUT/'source-snapshot';export.mkdir(exist_ok=True)
    manifest=[];rows=[];code_size=8.5;x=78;right=555
    for file_id,relative in enumerate(SOURCE_FILES,1):
        path=ROOT/relative;data=path.read_bytes();text=data.decode('utf-8-sig')
        if re.search(r'(?:tt[0-9a-f]{14,}|adunit-[0-9a-f]{6,}|AKIA[A-Z0-9]{16})',text,re.I):
            raise ValueError('Possible deployment identifier found in allowlisted file: '+relative)
        target=export/relative;target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(data)
        lines=text.splitlines();start=len(rows)
        for line_no,line in enumerate(lines,1):
            for continuation,segment in enumerate(fold_line(line,right-x,code_size)):
                rows.append({'file_id':file_id,'path':relative,'line':line_no,'continuation':continuation,'text':segment})
        manifest.append({'id':file_id,'path':relative,'bytes':len(data),'physical_lines':len(lines),
                         'printed_rows':len(rows)-start,'sha256':sha(data)})
    pages=[rows[i:i+50] for i in range(0,len(rows),50)];full_count=len(pages)
    chosen=list(range(full_count)) if full_count<60 else list(range(30))+list(range(full_count-30,full_count))
    target=OUT/(STEM+'_程序鉴别材料_内部草稿.pdf')
    c=canvas.Canvas(str(target),pagesize=A4,pageCompression=1)
    c.setTitle(NAME+' '+VERSION+' 程序鉴别材料 内部草稿');c.setAuthor('')
    c.setFont('CJK',24);c.drawString(55,751,NAME)
    c.setFont('CJK',17);c.drawString(55,714,VERSION+' 程序鉴别材料')
    c.setFont('CJK',12);c.drawString(55,677,'内部草稿 / 待权利人确认')
    notes=['本封面为内部审阅说明，不计入后续程序页码。',
        f'本次按真实源文件顺序汇集{len(manifest)}个文件，全量排版{full_count}页。',
        '程序正文交存方式草案：'+('不足60页，全部列出。' if full_count<60 else '连续前30页及连续后30页。'),
        '正文每页50个排版行，末页按实际剩余内容，不增加空代码凑页。',
        '长源行按版宽折行，左侧为文件编号:原始行号，+表示该行续排。',
        '页眉标明软件名称、版本和程序页码；原始源文件及SHA-256另附。',
        '空行来自源文件，未另增凑页；折行不是新增逻辑代码。',
        '本材料不含本地配置、真实广告ID、生成bundle或第三方库。',
        '权利人、开发日期、开发方式及申请表事实待本人确认；未填写登记号。',
        '提交前须核对受理端对行数、折行、封面及文件格式的具体要求。',
        '内部草稿应核对后再用于正式登记，不能视为已登记证明。']
    y=633
    for note in notes:c.setFont('CJK',10.5);c.drawString(55,y,note);y-=24
    y-=8;c.setFont('CJK',11);c.drawString(55,y,'文件编号与正文来源');y-=23
    for item in manifest:
        c.setFont('Mono',9);c.drawString(60,y,f'{item["id"]:02}  {item["path"]}')
        c.setFont('CJK',9);c.drawRightString(545,y,f'{item["physical_lines"]} 原始行');y-=19
    c.showPage();mapping=[]
    for output_page,original_page in enumerate(chosen,1):
        page=pages[original_page];c.setFillColorRGB(0,0,0);c.setFont('CJK',10)
        c.drawString(40,813,NAME+' '+VERSION+'  程序鉴别材料');c.drawRightString(555,813,str(output_page))
        first,last=page[0],page[-1];c.setFont('Mono',7.5)
        c.drawString(40,794,f'{first["path"]}:{first["line"]} - {last["path"]}:{last["line"]}')
        c.setLineWidth(.4);c.line(40,785,555,785);y=770
        for line_index,row in enumerate(page,1):
            prefix=f'{row["file_id"]:02}:{row["line"]:04}'+('+' if row['continuation'] else '')
            c.setFillColorRGB(.35,.35,.35);c.setFont('Mono',6.2);c.drawRightString(73,y,prefix)
            c.setFillColorRGB(0,0,0);draw_code(c,row['text'],x,y,code_size)
            mapping.append({**row,'pdf_page':output_page+1,'program_page':output_page,'printed_row':line_index});y-=14.2
        c.setFont('CJK',8);c.drawString(40,40,'内部草稿 / 待权利人确认')
        c.drawRightString(555,40,f'本页{len(page)}排版行 / 全量程序第{original_page+1}页');c.showPage()
    c.save()
    for item in manifest:
        actual=(ROOT/item['path']).read_text(encoding='utf-8-sig').splitlines();recovered={}
        for row in rows:
            if row['path']==item['path']:recovered.setdefault(row['line'],[]).append(row['text'])
        assert all(''.join(recovered[i+1])==line.expandtabs(4) for i,line in enumerate(actual))
    metadata={'software':NAME,'version':VERSION,'status':'内部草稿 / 待权利人确认',
        'snapshot_created_utc':datetime.now(timezone.utc).isoformat(),'source_files':manifest,
        'original_lines':sum(i['physical_lines'] for i in manifest),'printed_rows':len(rows),
        'full_program_pages':full_count,'included_program_pages':len(chosen),'cover_pages':1,
        'page_size':'A4','program_rows_per_page':50,'last_page_rows':len(pages[-1]),
        'folding_verified_reversible':True,'excluded':['config.local.json','build/**','web/game.bundle.js','tests/**','node_modules/**'],
        'official_reference':SOURCES_URL}
    (OUT/'source-manifest.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with (OUT/'source-page-map.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(mapping[0]));writer.writeheader();writer.writerows(mapping)
    hashes=''.join(f'{i["sha256"]}  source-snapshot/{i["path"]}\n' for i in manifest)
    (OUT/'SHA256SUMS.txt').write_text(hashes,encoding='utf-8');return metadata

def main():
    p=argparse.ArgumentParser();p.add_argument('--frozen',action='store_true')
    p.add_argument('--screenshots',type=Path);p.add_argument('--manual-only',action='store_true')
    args=p.parse_args();OUT.mkdir(parents=True,exist_ok=True)
    if not args.manual_only and not args.frozen:raise SystemExit('Source export requires explicit --frozen.')
    shots=json.loads(args.screenshots.read_text(encoding='utf-8')) if args.screenshots else []
    manual=make_manual(shots);summary={'manual_docx':str(manual),'screenshots':len(shots)}
    if not args.manual_only:summary.update(make_sources())
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
